"""Shop Billing App — daily billing for sanitary, paints, plumbing, electrical shops.

Run:  python app.py   then open http://localhost:5000
Default login: admin / admin123  (change it from the Users page)

Works on desktop and tablet: open http://<computer-ip>:5000 from the tablet
browser on the same Wi-Fi.
"""

import io
import os
import shutil
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (Flask, Response, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(APP_DIR, "shop.db")

app = Flask(__name__)
app.secret_key = os.environ.get("SHOP_SECRET", "change-this-secret-key")

CATEGORIES = ["Sanitary", "Paints", "Plumbing", "Electrical", "Hardware", "Other"]
GST_RATES = [0, 5, 12, 18, 28]

SEED_PRODUCTS = [
    # (category, name, description, unit, price, purchase_price, gst_rate, stock, low_stock)
    ("Sanitary", "Wash Basin", "Ceramic wash basin, white", "pcs", 1450, 1100, 18, 10, 3),
    ("Sanitary", "Western Toilet Seat", "Floor-mounted EWC with flush tank", "pcs", 5200, 4100, 18, 5, 2),
    ("Sanitary", "Health Faucet", "ABS body health faucet with hose", "pcs", 350, 240, 18, 25, 5),
    ("Sanitary", "Bathroom Mirror", "18x24 inch frameless mirror", "pcs", 650, 450, 18, 12, 3),
    ("Paints", "Emulsion Paint 20L", "Interior emulsion, white base", "bucket", 3800, 3100, 18, 8, 2),
    ("Paints", "Enamel Paint 1L", "Oil-based enamel, gloss finish", "tin", 320, 240, 18, 30, 6),
    ("Paints", "Primer 10L", "Wall primer, water-based", "bucket", 1600, 1250, 18, 10, 2),
    ("Paints", "Paint Brush 4 inch", "Bristle brush for oil/water paints", "pcs", 90, 55, 18, 40, 10),
    ("Plumbing", "PVC Pipe 1 inch", "3-metre length, ISI mark", "pcs", 210, 150, 18, 60, 15),
    ("Plumbing", "CPVC Pipe 3/4 inch", "3-metre length, hot/cold water", "pcs", 260, 190, 18, 50, 15),
    ("Plumbing", "Ball Valve 1/2 inch", "Brass ball valve", "pcs", 180, 120, 18, 35, 8),
    ("Plumbing", "Teflon Tape", "Thread seal tape roll", "pcs", 15, 8, 12, 100, 25),
    ("Electrical", "Wire 1.5 sqmm 90m", "FR copper wire coil", "coil", 1750, 1400, 18, 15, 4),
    ("Electrical", "Switch 6A", "Modular switch, white", "pcs", 45, 28, 18, 80, 20),
    ("Electrical", "LED Bulb 9W", "B22 cool white LED bulb", "pcs", 110, 75, 12, 60, 15),
    ("Electrical", "MCB 16A", "Single-pole miniature circuit breaker", "pcs", 240, 170, 18, 20, 5),
    ("Hardware", "Door Hinge 4 inch", "Stainless steel butt hinge", "pcs", 60, 38, 18, 50, 12),
    ("Hardware", "Screws Box 1 inch", "100-piece box, zinc plated", "box", 120, 80, 18, 30, 8),
]


# ---------------------------------------------------------------- database

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def ensure_column(db, table, col, ddl):
    cols = [r[1] for r in db.execute(f"PRAGMA table_info({table})")]
    if col not in cols:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {col} {ddl}")


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            username      TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            role          TEXT NOT NULL DEFAULT 'staff',   -- admin | staff
            active        INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS customers (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            name       TEXT NOT NULL,
            mobile     TEXT DEFAULT '',
            address    TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            active     INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS products (
            id             INTEGER PRIMARY KEY AUTOINCREMENT,
            category       TEXT NOT NULL,
            name           TEXT NOT NULL,
            description    TEXT DEFAULT '',
            unit           TEXT DEFAULT 'pcs',
            price          REAL NOT NULL DEFAULT 0,
            purchase_price REAL NOT NULL DEFAULT 0,
            gst_rate       REAL NOT NULL DEFAULT 0,
            stock          REAL NOT NULL DEFAULT 0,
            low_stock      REAL NOT NULL DEFAULT 5,
            active         INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS bills (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            bill_no       TEXT NOT NULL UNIQUE,
            customer_id   INTEGER REFERENCES customers(id),
            customer_name TEXT NOT NULL,
            phone         TEXT DEFAULT '',
            bill_date     TEXT NOT NULL,           -- YYYY-MM-DD
            created_at    TEXT NOT NULL,
            created_by    TEXT DEFAULT '',
            subtotal      REAL NOT NULL DEFAULT 0, -- sum of qty*price
            item_discount REAL NOT NULL DEFAULT 0,
            discount      REAL NOT NULL DEFAULT 0, -- extra overall discount
            gst_amount    REAL NOT NULL DEFAULT 0,
            total         REAL NOT NULL DEFAULT 0,
            payment_mode  TEXT NOT NULL DEFAULT 'Cash'
        );
        CREATE TABLE IF NOT EXISTS bill_items (
            id             INTEGER PRIMARY KEY AUTOINCREMENT,
            bill_id        INTEGER NOT NULL REFERENCES bills(id) ON DELETE CASCADE,
            product_id     INTEGER,
            category       TEXT NOT NULL,
            product_name   TEXT NOT NULL,
            description    TEXT DEFAULT '',
            unit           TEXT DEFAULT 'pcs',
            qty            REAL NOT NULL,
            price          REAL NOT NULL,
            purchase_price REAL NOT NULL DEFAULT 0,
            discount       REAL NOT NULL DEFAULT 0,
            gst_rate       REAL NOT NULL DEFAULT 0,
            gst_amount     REAL NOT NULL DEFAULT 0,
            amount         REAL NOT NULL
        );
        CREATE TABLE IF NOT EXISTS stock_moves (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            product_id INTEGER NOT NULL REFERENCES products(id),
            change     REAL NOT NULL,      -- +in / -out
            reason     TEXT NOT NULL,      -- Stock In | Sale | Adjustment
            ref        TEXT DEFAULT '',
            note       TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            created_by TEXT DEFAULT ''
        );
        """
    )
    # migrate a v1 database (pre-GST schema) in place
    for col, ddl in [("purchase_price", "REAL NOT NULL DEFAULT 0"),
                     ("gst_rate", "REAL NOT NULL DEFAULT 0"),
                     ("stock", "REAL NOT NULL DEFAULT 0"),
                     ("low_stock", "REAL NOT NULL DEFAULT 5")]:
        ensure_column(db, "products", col, ddl)
    for col, ddl in [("customer_id", "INTEGER"), ("created_by", "TEXT DEFAULT ''"),
                     ("item_discount", "REAL NOT NULL DEFAULT 0"),
                     ("gst_amount", "REAL NOT NULL DEFAULT 0")]:
        ensure_column(db, "bills", col, ddl)
    for col, ddl in [("product_id", "INTEGER"),
                     ("purchase_price", "REAL NOT NULL DEFAULT 0"),
                     ("discount", "REAL NOT NULL DEFAULT 0"),
                     ("gst_rate", "REAL NOT NULL DEFAULT 0"),
                     ("gst_amount", "REAL NOT NULL DEFAULT 0")]:
        ensure_column(db, "bill_items", col, ddl)

    if db.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 0:
        db.executemany(
            "INSERT INTO products (category, name, description, unit, price,"
            " purchase_price, gst_rate, stock, low_stock) VALUES (?,?,?,?,?,?,?,?,?)",
            SEED_PRODUCTS,
        )
    if db.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        db.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?,?,?)",
            ("admin", generate_password_hash("admin123"), "admin"),
        )
    db.commit()
    db.close()


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def next_bill_no(db):
    today = date.today().strftime("%Y%m%d")
    prefix = f"INV-{today}-"
    row = db.execute(
        "SELECT bill_no FROM bills WHERE bill_no LIKE ? ORDER BY id DESC LIMIT 1",
        (prefix + "%",),
    ).fetchone()
    seq = int(row["bill_no"].rsplit("-", 1)[1]) + 1 if row else 1
    return f"{prefix}{seq:03d}"


# ---------------------------------------------------------------- auth

def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            return render_template("message.html",
                                   title="Not allowed",
                                   text="Only an admin user can open this page."), 403
        return fn(*args, **kwargs)
    return wrapper


@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        u = (request.form.get("username") or "").strip()
        p = request.form.get("password") or ""
        row = get_db().execute(
            "SELECT * FROM users WHERE username = ? AND active = 1", (u,)
        ).fetchone()
        if row and check_password_hash(row["password_hash"], p):
            session["user"] = row["username"]
            session["role"] = row["role"]
            return redirect(request.args.get("next") or url_for("dashboard"))
        error = "Wrong username or password"
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------- dashboard

@app.route("/")
@login_required
def dashboard():
    db = get_db()
    today = date.today()
    iso = today.isoformat()

    def sales_between(start, end):
        return db.execute(
            "SELECT COUNT(*) AS bills, COALESCE(SUM(total),0) AS sales"
            " FROM bills WHERE bill_date BETWEEN ? AND ?", (start, end)
        ).fetchone()

    week_start = (today - timedelta(days=today.weekday())).isoformat()
    tiles = {
        "Today": sales_between(iso, iso),
        "This Week": sales_between(week_start, iso),
        "This Month": sales_between(today.replace(day=1).isoformat(), iso),
        "This Year": sales_between(today.replace(month=1, day=1).isoformat(), iso),
    }

    # last-7-days trend for the bar chart
    days = [(today - timedelta(days=i)) for i in range(6, -1, -1)]
    sales_by_day = dict(db.execute(
        "SELECT bill_date, SUM(total) FROM bills WHERE bill_date >= ?"
        " GROUP BY bill_date", (days[0].isoformat(),)
    ).fetchall())
    trend = [{"label": d.strftime("%a"), "date": d.isoformat(),
              "value": round(sales_by_day.get(d.isoformat(), 0) or 0, 2)}
             for d in days]
    trend_max = max([t["value"] for t in trend] + [1])

    low_stock = db.execute(
        "SELECT * FROM products WHERE active = 1 AND stock <= low_stock"
        " ORDER BY (stock - low_stock) LIMIT 10"
    ).fetchall()
    recent = db.execute("SELECT * FROM bills ORDER BY id DESC LIMIT 5").fetchall()
    top_products = db.execute(
        "SELECT i.product_name, SUM(i.qty) AS qty, SUM(i.amount) AS sales"
        " FROM bill_items i JOIN bills b ON b.id = i.bill_id"
        " WHERE b.bill_date >= ? GROUP BY i.product_name"
        " ORDER BY sales DESC LIMIT 5",
        (today.replace(day=1).isoformat(),),
    ).fetchall()
    return render_template(
        "dashboard.html", tiles=tiles, trend=trend, trend_max=trend_max,
        low_stock=low_stock, recent=recent, top_products=top_products,
    )


# ---------------------------------------------------------------- customers

@app.route("/customers")
@login_required
def customers():
    q = (request.args.get("q") or "").strip()
    sql = ("SELECT c.*, COUNT(b.id) AS bills, COALESCE(SUM(b.total),0) AS sales"
           " FROM customers c LEFT JOIN bills b ON b.customer_id = c.id"
           " WHERE c.active = 1")
    args = []
    if q:
        sql += " AND (c.name LIKE ? OR c.mobile LIKE ?)"
        args = [f"%{q}%"] * 2
    rows = get_db().execute(sql + " GROUP BY c.id ORDER BY c.name", args).fetchall()
    return render_template("customers.html", customers=rows, q=q)


@app.route("/customers/add", methods=["POST"])
@login_required
def add_customer():
    f = request.form
    name = (f.get("name") or "").strip()
    if name:
        db = get_db()
        db.execute(
            "INSERT INTO customers (name, mobile, address, created_at) VALUES (?,?,?,?)",
            (name, (f.get("mobile") or "").strip(), (f.get("address") or "").strip(), now()),
        )
        db.commit()
    return redirect(url_for("customers"))


@app.route("/customers/<int:cid>/edit", methods=["POST"])
@login_required
def edit_customer(cid):
    f = request.form
    db = get_db()
    db.execute(
        "UPDATE customers SET name = ?, mobile = ?, address = ? WHERE id = ?",
        ((f.get("name") or "").strip(), (f.get("mobile") or "").strip(),
         (f.get("address") or "").strip(), cid),
    )
    db.commit()
    return redirect(url_for("customers"))


@app.route("/customers/<int:cid>/delete", methods=["POST"])
@login_required
def delete_customer(cid):
    db = get_db()
    db.execute("UPDATE customers SET active = 0 WHERE id = ?", (cid,))
    db.commit()
    return redirect(url_for("customers"))


def find_or_create_customer(db, name, mobile, address=""):
    row = None
    if mobile:
        row = db.execute(
            "SELECT id FROM customers WHERE mobile = ? AND active = 1", (mobile,)
        ).fetchone()
    if not row:
        row = db.execute(
            "SELECT id FROM customers WHERE name = ? COLLATE NOCASE AND active = 1",
            (name,),
        ).fetchone()
    if row:
        return row["id"]
    cur = db.execute(
        "INSERT INTO customers (name, mobile, address, created_at) VALUES (?,?,?,?)",
        (name, mobile, address, now()),
    )
    return cur.lastrowid


# ---------------------------------------------------------------- billing

@app.route("/billing")
@login_required
def new_bill():
    db = get_db()
    products = db.execute(
        "SELECT * FROM products WHERE active = 1 ORDER BY category, name"
    ).fetchall()
    custs = db.execute(
        "SELECT * FROM customers WHERE active = 1 ORDER BY name LIMIT 500"
    ).fetchall()
    return render_template(
        "billing.html", categories=CATEGORIES, gst_rates=GST_RATES,
        products=[dict(p) for p in products],
        customers=[dict(c) for c in custs],
        today=date.today().isoformat(),
    )


@app.route("/api/bills", methods=["POST"])
@login_required
def save_bill():
    data = request.get_json(force=True)
    name = (data.get("customer_name") or "").strip()
    items = data.get("items") or []
    if not name:
        return jsonify(error="Customer name is required"), 400
    if not items:
        return jsonify(error="Add at least one item"), 400

    db = get_db()
    clean, subtotal, item_disc, gst_total = [], 0.0, 0.0, 0.0
    for it in items:
        try:
            qty = float(it.get("qty") or 0)
            price = float(it.get("price") or 0)
            disc = max(0.0, float(it.get("discount") or 0))
            gst_rate = max(0.0, float(it.get("gst_rate") or 0))
        except (TypeError, ValueError):
            return jsonify(error="Qty, price, discount and GST must be numbers"), 400
        pname = (it.get("product_name") or "").strip()
        if not pname or qty <= 0:
            return jsonify(error="Each item needs a product and a quantity above 0"), 400

        pid, cost = None, 0.0
        if it.get("product_id"):
            prod = db.execute(
                "SELECT * FROM products WHERE id = ?", (it["product_id"],)
            ).fetchone()
            if prod:
                pid, cost = prod["id"], prod["purchase_price"]

        base = qty * price - disc
        gst_amt = round(base * gst_rate / 100, 2)
        amount = round(base + gst_amt, 2)
        subtotal += qty * price
        item_disc += disc
        gst_total += gst_amt
        clean.append((pid, it.get("category") or "Other", pname,
                      (it.get("description") or "").strip(),
                      it.get("unit") or "pcs", qty, price, cost,
                      disc, gst_rate, gst_amt, amount))

    try:
        discount = max(0.0, float(data.get("discount") or 0))
    except (TypeError, ValueError):
        discount = 0.0
    total = round(subtotal - item_disc + gst_total - discount, 2)

    cust_id = find_or_create_customer(
        db, name, (data.get("phone") or "").strip(), (data.get("address") or "").strip()
    )
    bill_no = next_bill_no(db)
    cur = db.execute(
        "INSERT INTO bills (bill_no, customer_id, customer_name, phone, bill_date,"
        " created_at, created_by, subtotal, item_discount, discount, gst_amount,"
        " total, payment_mode) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (bill_no, cust_id, name, (data.get("phone") or "").strip(),
         data.get("bill_date") or date.today().isoformat(), now(),
         session.get("user", ""), round(subtotal, 2), round(item_disc, 2),
         round(discount, 2), round(gst_total, 2), total,
         data.get("payment_mode") or "Cash"),
    )
    bill_id = cur.lastrowid
    db.executemany(
        "INSERT INTO bill_items (bill_id, product_id, category, product_name,"
        " description, unit, qty, price, purchase_price, discount, gst_rate,"
        " gst_amount, amount) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        [(bill_id,) + item for item in clean],
    )
    # stock out for every catalogue item on the bill
    for item in clean:
        pid, qty = item[0], item[5]
        if pid:
            db.execute("UPDATE products SET stock = stock - ? WHERE id = ?", (qty, pid))
            db.execute(
                "INSERT INTO stock_moves (product_id, change, reason, ref,"
                " created_at, created_by) VALUES (?,?,?,?,?,?)",
                (pid, -qty, "Sale", bill_no, now(), session.get("user", "")),
            )
    db.commit()
    return jsonify(ok=True, bill_id=bill_id, bill_no=bill_no)


@app.route("/bills")
@login_required
def bills():
    db = get_db()
    q = (request.args.get("q") or "").strip()
    day = request.args.get("date") or ""
    sql = "SELECT * FROM bills WHERE 1=1"
    args = []
    if q:
        sql += " AND (customer_name LIKE ? OR phone LIKE ? OR bill_no LIKE ?)"
        args += [f"%{q}%"] * 3
    if day:
        sql += " AND bill_date = ?"
        args.append(day)
    rows = db.execute(sql + " ORDER BY id DESC LIMIT 300", args).fetchall()
    total = sum(r["total"] for r in rows)
    return render_template("bills.html", bills=rows, q=q, day=day, total=total)


@app.route("/bills/<int:bill_id>")
@login_required
def bill_view(bill_id):
    db = get_db()
    bill = db.execute("SELECT * FROM bills WHERE id = ?", (bill_id,)).fetchone()
    if not bill:
        return "Bill not found", 404
    items = db.execute(
        "SELECT * FROM bill_items WHERE bill_id = ? ORDER BY id", (bill_id,)
    ).fetchall()
    customer = None
    if bill["customer_id"]:
        customer = db.execute(
            "SELECT * FROM customers WHERE id = ?", (bill["customer_id"],)
        ).fetchone()
    return render_template("bill_view.html", bill=bill, items=items, customer=customer)


# ---------------------------------------------------------------- products

@app.route("/products")
@login_required
def products():
    rows = get_db().execute(
        "SELECT * FROM products WHERE active = 1 ORDER BY category, name"
    ).fetchall()
    return render_template("products.html", products=rows,
                           categories=CATEGORIES, gst_rates=GST_RATES)


def product_form_values(f):
    def num(key, default=0.0):
        try:
            return float(f.get(key) or default)
        except ValueError:
            return default
    return ((f.get("category") or "Other"), (f.get("name") or "").strip(),
            (f.get("description") or "").strip(), (f.get("unit") or "pcs").strip(),
            num("price"), num("purchase_price"), num("gst_rate"),
            num("stock"), num("low_stock", 5))


@app.route("/products/add", methods=["POST"])
@login_required
def add_product():
    vals = product_form_values(request.form)
    if vals[1]:
        db = get_db()
        cur = db.execute(
            "INSERT INTO products (category, name, description, unit, price,"
            " purchase_price, gst_rate, stock, low_stock) VALUES (?,?,?,?,?,?,?,?,?)",
            vals,
        )
        if vals[7] > 0:
            db.execute(
                "INSERT INTO stock_moves (product_id, change, reason, note,"
                " created_at, created_by) VALUES (?,?,?,?,?,?)",
                (cur.lastrowid, vals[7], "Stock In", "Opening stock", now(),
                 session.get("user", "")),
            )
        db.commit()
    return redirect(url_for("products"))


@app.route("/products/<int:pid>/edit", methods=["POST"])
@login_required
def edit_product(pid):
    vals = product_form_values(request.form)
    if vals[1]:
        db = get_db()
        db.execute(
            "UPDATE products SET category=?, name=?, description=?, unit=?,"
            " price=?, purchase_price=?, gst_rate=?, stock=?, low_stock=? WHERE id=?",
            vals + (pid,),
        )
        db.commit()
    return redirect(url_for("products"))


@app.route("/products/<int:pid>/delete", methods=["POST"])
@login_required
def delete_product(pid):
    db = get_db()
    db.execute("UPDATE products SET active = 0 WHERE id = ?", (pid,))
    db.commit()
    return redirect(url_for("products"))


# ---------------------------------------------------------------- inventory

@app.route("/inventory")
@login_required
def inventory():
    db = get_db()
    prods = db.execute(
        "SELECT * FROM products WHERE active = 1 ORDER BY category, name"
    ).fetchall()
    low = [p for p in prods if p["stock"] <= p["low_stock"]]
    moves = db.execute(
        "SELECT m.*, p.name AS product_name, p.unit FROM stock_moves m"
        " JOIN products p ON p.id = m.product_id"
        " ORDER BY m.id DESC LIMIT 100"
    ).fetchall()
    return render_template("inventory.html", products=prods, low=low, moves=moves)


@app.route("/inventory/move", methods=["POST"])
@login_required
def stock_move():
    f = request.form
    try:
        pid = int(f.get("product_id") or 0)
        qty = float(f.get("qty") or 0)
    except ValueError:
        pid, qty = 0, 0
    direction = f.get("direction") or "in"
    if pid and qty > 0:
        change = qty if direction == "in" else -qty
        reason = "Stock In" if direction == "in" else "Adjustment"
        db = get_db()
        db.execute("UPDATE products SET stock = stock + ? WHERE id = ?", (change, pid))
        db.execute(
            "INSERT INTO stock_moves (product_id, change, reason, note,"
            " created_at, created_by) VALUES (?,?,?,?,?,?)",
            (pid, change, reason, (f.get("note") or "").strip(), now(),
             session.get("user", "")),
        )
        db.commit()
    return redirect(url_for("inventory"))


# ---------------------------------------------------------------- reports

PERIOD_SQL = {
    "daily": "bill_date",
    "weekly": "strftime('%Y', bill_date) || '-W' || strftime('%W', bill_date)",
    "monthly": "strftime('%Y-%m', bill_date)",
    "quarterly": "strftime('%Y', bill_date) || '-Q' || ((strftime('%m', bill_date) + 2) / 3)",
    "yearly": "strftime('%Y', bill_date)",
}
PERIODS = list(PERIOD_SQL)
REPORT_TYPES = ["sales", "category", "product", "customer", "profit"]


def default_range(period):
    today = date.today()
    if period == "daily":
        start = today - timedelta(days=30)
    elif period == "weekly":
        start = today - timedelta(weeks=12)
    elif period == "monthly":
        start = today.replace(month=1, day=1)
    else:
        start = today.replace(year=today.year - 4, month=1, day=1)
    return start.isoformat(), today.isoformat()


def report_data(db, rtype, period, start, end):
    """Return (headers, rows) for a report; rows are plain lists."""
    bucket = PERIOD_SQL[period]
    rng = (start, end)
    if rtype == "sales":
        rows = db.execute(
            f"SELECT {bucket} AS period, COUNT(*) AS bills, SUM(subtotal),"
            f" SUM(item_discount + discount), SUM(gst_amount), SUM(total)"
            f" FROM bills WHERE bill_date BETWEEN ? AND ?"
            f" GROUP BY period ORDER BY period DESC", rng).fetchall()
        return (["Period", "Bills", "Subtotal", "Discount", "GST", "Sales"],
                [list(r) for r in rows])
    if rtype == "category":
        rows = db.execute(
            "SELECT i.category, COUNT(DISTINCT b.id), SUM(i.qty), SUM(i.amount)"
            " FROM bill_items i JOIN bills b ON b.id = i.bill_id"
            " WHERE b.bill_date BETWEEN ? AND ?"
            " GROUP BY i.category ORDER BY SUM(i.amount) DESC", rng).fetchall()
        return (["Category", "Bills", "Qty", "Sales"], [list(r) for r in rows])
    if rtype == "product":
        rows = db.execute(
            "SELECT i.product_name, i.category, SUM(i.qty), SUM(i.amount)"
            " FROM bill_items i JOIN bills b ON b.id = i.bill_id"
            " WHERE b.bill_date BETWEEN ? AND ?"
            " GROUP BY i.product_name, i.category"
            " ORDER BY SUM(i.amount) DESC LIMIT 100", rng).fetchall()
        return (["Product", "Category", "Qty", "Sales"], [list(r) for r in rows])
    if rtype == "customer":
        rows = db.execute(
            "SELECT customer_name, phone, COUNT(*), SUM(total)"
            " FROM bills WHERE bill_date BETWEEN ? AND ?"
            " GROUP BY customer_name, phone"
            " ORDER BY SUM(total) DESC LIMIT 100", rng).fetchall()
        return (["Customer", "Phone", "Bills", "Sales"], [list(r) for r in rows])
    # profit: revenue before GST minus cost, per period bucket
    rows = db.execute(
        f"SELECT {bucket} AS period,"
        f" SUM(i.qty * i.price - i.discount) AS revenue,"
        f" SUM(i.qty * i.purchase_price) AS cost,"
        f" SUM(i.qty * i.price - i.discount - i.qty * i.purchase_price) AS profit"
        f" FROM bill_items i JOIN bills b ON b.id = i.bill_id"
        f" WHERE b.bill_date BETWEEN ? AND ?"
        f" GROUP BY period ORDER BY period DESC", rng).fetchall()
    out = []
    for r in rows:
        margin = (r["profit"] / r["revenue"] * 100) if r["revenue"] else 0
        out.append([r["period"], r["revenue"], r["cost"], r["profit"],
                    round(margin, 1)])
    return (["Period", "Revenue (excl GST)", "Cost", "Profit", "Margin %"], out)


def report_params():
    rtype = request.args.get("type", "sales")
    if rtype not in REPORT_TYPES:
        rtype = "sales"
    period = request.args.get("period", "daily")
    if period not in PERIOD_SQL:
        period = "daily"
    d_start, d_end = default_range(period)
    return rtype, period, (request.args.get("start") or d_start), (request.args.get("end") or d_end)


@app.route("/reports")
@login_required
def reports():
    rtype, period, start, end = report_params()
    headers, rows = report_data(get_db(), rtype, period, start, end)
    return render_template(
        "reports.html", rtype=rtype, period=period, start=start, end=end,
        headers=headers, rows=rows, periods=PERIODS, report_types=REPORT_TYPES,
    )


@app.route("/reports/export.xlsx")
@login_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rtype, period, start, end = report_params()
    headers, rows = report_data(get_db(), rtype, period, start, end)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{rtype}-{period}"[:31]
    ws.append([f"{rtype.capitalize()} report ({period}) {start} to {end}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([round(v, 2) if isinstance(v, float) else v for v in r])
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = max(14, len(str(h)) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f"{rtype}_{period}_{start}_to_{end}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/reports/export.csv")
@login_required
def export_csv():
    import csv as _csv
    rtype, period, start, end = report_params()
    headers, rows = report_data(get_db(), rtype, period, start, end)
    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow([round(v, 2) if isinstance(v, float) else v for v in r])
    return Response(
        buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition":
                 f"attachment; filename={rtype}_{period}_{start}_to_{end}.csv"},
    )


# ---------------------------------------------------------------- users (admin)

@app.route("/users")
@admin_required
def users():
    rows = get_db().execute("SELECT * FROM users ORDER BY username").fetchall()
    return render_template("users.html", users=rows)


@app.route("/users/add", methods=["POST"])
@admin_required
def add_user():
    f = request.form
    u = (f.get("username") or "").strip()
    p = f.get("password") or ""
    if u and len(p) >= 4:
        db = get_db()
        try:
            db.execute(
                "INSERT INTO users (username, password_hash, role) VALUES (?,?,?)",
                (u, generate_password_hash(p),
                 "admin" if f.get("role") == "admin" else "staff"),
            )
            db.commit()
        except sqlite3.IntegrityError:
            pass  # username already exists
    return redirect(url_for("users"))


@app.route("/users/<int:uid>/toggle", methods=["POST"])
@admin_required
def toggle_user(uid):
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
    if row and row["username"] != session.get("user"):
        db.execute("UPDATE users SET active = 1 - active WHERE id = ?", (uid,))
        db.commit()
    return redirect(url_for("users"))


@app.route("/users/<int:uid>/password", methods=["POST"])
@admin_required
def reset_password(uid):
    p = request.form.get("password") or ""
    if len(p) >= 4:
        db = get_db()
        db.execute("UPDATE users SET password_hash = ? WHERE id = ?",
                   (generate_password_hash(p), uid))
        db.commit()
    return redirect(url_for("users"))


# ---------------------------------------------------------------- backup

@app.route("/backup")
@admin_required
def backup():
    size = os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0
    return render_template("backup.html", size_kb=round(size / 1024, 1))


@app.route("/backup/download")
@admin_required
def backup_download():
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return send_file(DB_PATH, as_attachment=True,
                     download_name=f"shop-backup-{stamp}.db")


@app.route("/backup/restore", methods=["POST"])
@admin_required
def backup_restore():
    f = request.files.get("file")
    if not f:
        return render_template("message.html", title="Restore failed",
                               text="No file was uploaded."), 400
    head = f.stream.read(16)
    f.stream.seek(0)
    if head != b"SQLite format 3\x00":
        return render_template("message.html", title="Restore failed",
                               text="That file is not a valid backup (.db) file."), 400
    # keep the current database aside before replacing it
    if os.path.exists(DB_PATH):
        shutil.copy2(DB_PATH, DB_PATH + ".before-restore")
    f.save(DB_PATH)
    init_db()  # apply any missing migrations to the restored file
    return render_template("message.html", title="Restore complete",
                           text="The backup was restored successfully.")


if __name__ == "__main__":
    init_db()
    # host 0.0.0.0 lets tablets/phones on the same Wi-Fi open the app
    app.run(host="0.0.0.0", port=5000, debug=False)
