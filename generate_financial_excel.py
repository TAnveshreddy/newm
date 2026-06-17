"""Generate FinancialAnalysis_Data.xlsx — standalone data file matching the PBIT."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MONTHS = [
    ("Jan", 1, 1650000, 520000, 480000, 78000),
    ("Feb", 2, 1420000, 460000, 410000, 65000),
    ("Mar", 3, 1380000, 510000, 380000, 70000),
    ("Apr", 4, 1200000, 550000, 360000, 60000),
    ("May", 5, 1510000, 490000, 430000, 71000),
    ("Jun", 6, 1480000, 480000, 440000, 68000),
    ("Jul", 7, 1350000, 510000, 400000, 63000),
    ("Aug", 8, 1290000, 490000, 390000, 61000),
    ("Sep", 9, 1100000, 580000, 360000, 70000),
    ("Oct",10, 1420000, 500000, 400000, 68000),
    ("Nov",11, 1250000, 540000, 390000, 65000),
    ("Dec",12, 1560000, 450000, 450000, 62000),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "FinancialData"

DARK  = "1F2232"
GOLD  = "F5A623"
MID   = "2D3250"
WHITE = "FFFFFF"
GREY  = "C8C8C8"

headers = ["Month","MonthNum","Year","Revenue","COGS","GrossProfit",
           "Opex","EBIT","InterestTax","NetProfit","NetProfitMargin","BreakEven","Expenses"]

hdr_fill = PatternFill("solid", fgColor=GOLD)
row_fill  = PatternFill("solid", fgColor=MID)
alt_fill  = PatternFill("solid", fgColor=DARK)
thin = Side(style="thin", color="444444")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for ci, h in enumerate(headers, 1):
    cell = ws.cell(1, ci, h)
    cell.font = Font(bold=True, color=DARK, name="Segoe UI")
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for ri, (mn, mnum, rev, cogs, opex, itax) in enumerate(MONTHS, 2):
    gp   = rev - cogs
    ebit = gp - opex
    np_  = ebit - itax
    npm  = round(np_ / rev * 100, 1)
    be   = cogs + opex + itax
    exp  = cogs + opex + itax
    row_data = [mn, mnum, 2023, rev, cogs, gp, opex, ebit, itax, np_, npm, be, exp]
    fill = row_fill if ri % 2 == 0 else alt_fill
    for ci, val in enumerate(row_data, 1):
        cell = ws.cell(ri, ci, val)
        cell.font = Font(color=WHITE if ci > 3 else GREY, name="Segoe UI")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="right" if ci > 1 else "left")
        cell.border = border
        if ci in (4,5,6,7,8,9,10,12,13):
            cell.number_format = '$#,##0'
        elif ci == 11:
            cell.number_format = '0.0"%"'

# Totals row
tr = len(MONTHS) + 2
totals = ["TOTAL", "", 2023]
all_rows = [(rev,cogs,rev-cogs,opex,rev-cogs-opex,itax,
             rev-cogs-opex-itax,0,cogs+opex+itax,cogs+opex+itax)
            for _,_,rev,cogs,opex,itax in MONTHS]
totals += [sum(r[i] for r in all_rows) for i in range(10)]

tot_fill = PatternFill("solid", fgColor=GOLD)
for ci, val in enumerate(totals[:13], 1):
    cell = ws.cell(tr, ci, val if val != 0 else "")
    cell.font = Font(bold=True, color=DARK, name="Segoe UI")
    cell.fill = tot_fill
    cell.border = border
    cell.alignment = Alignment(horizontal="right" if ci > 1 else "left")
    if ci in (4,5,6,7,8,9,10,12,13):
        cell.number_format = '$#,##0'

# Column widths
widths = [8,9,6,12,12,14,10,12,13,12,14,12,12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 25

# Summary sheet
ws2 = wb.create_sheet("Summary")
ws2.sheet_view.showGridLines = False
ws2["A1"] = "Financial Analysis — 2023 Summary"
ws2["A1"].font = Font(bold=True, size=14, color=GOLD, name="Segoe UI")
ws2["A1"].fill = PatternFill("solid", fgColor=DARK)

kpis = [
    ("Total Revenue",      sum(r[0] for r in all_rows), "$#,##0.0,,\"M\""),
    ("Total Expenses",     sum(r[9] for r in all_rows), "$#,##0.0,,\"M\""),
    ("Total Gross Profit", sum(r[2] for r in all_rows), "$#,##0.0,,\"M\""),
    ("Total EBIT",         sum(r[4] for r in all_rows), "$#,##0.0,,\"M\""),
    ("Total Net Profit",   sum(r[6] for r in all_rows), "$#,##0.0,,\"M\""),
]

for i, (label, val, fmt) in enumerate(kpis, 3):
    ws2.cell(i, 1, label).font = Font(bold=True, color=GREY, name="Segoe UI")
    ws2.cell(i, 1).fill = PatternFill("solid", fgColor=MID)
    c = ws2.cell(i, 2, val)
    c.font = Font(bold=True, color=GOLD, name="Segoe UI", size=13)
    c.fill = PatternFill("solid", fgColor=MID)
    c.number_format = fmt
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 14

out = "/home/user/newm/FinancialAnalysis_Data.xlsx"
wb.save(out)
print(f"✓ Written: {out}")
